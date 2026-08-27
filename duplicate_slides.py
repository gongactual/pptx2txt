# /// script
# requires-python = ">=3.12"
# dependencies = [
#     "python-pptx",
# ]
# ///

# ================
# DUPLICATE_SLIDES
# ================
#
# VERSION
# =======
# 1.0.0 | 26 Aug 2026
#
# COPYRIGHT STATEMENT
# ===================
# This script was created by Oxzeon Limited (oxzeon.com) with assistance from Claude Code.
# No copyright is asserted - please use/modify/redistribute as you wish.
# CAUTION: Use entirely at your own risk.
# Contact: alan@oxzeon.com
#
# OVERVIEW
# ========
# This script opens a .pptx file, scans every slide for a shape whose text is wrapped in
# double square brackets (e.g. "[[Divider]]"), and builds an index mapping each such label to
# the slide that carries it. It also reads a JSON file (e.g. demo.json) containing a "slides"
# array, where each entry has a "layout" value naming one of these labels; the ordered list of
# "layout" values (duplicates allowed) becomes the list of slides to duplicate. It duplicates
# the corresponding slides - including independent copies of their notes and any chart-related
# content, so edits to a duplicate never affect the original - and prepends the duplicates to
# the front of the deck, in the order given by the JSON file. The result is saved to a new
# .pptx file. This script ignores Any "content" field on a JSON slide entry.
#
# Slide duplication is done by unzipping the .pptx (which is just a zip archive of XML parts)
# and manipulating the relevant XML/relationship files directly, since python-pptx has no
# built-in slide-duplication feature. Images, media and slide layouts are left shared with the
# original slide (matching PowerPoint's own "Duplicate Slide" behaviour for pictures).
#
# REQUIREMENTS
# ============
# - Python 3.12+
# - Library: python-pptx
#
# USAGE
# =====
# python duplicate_slides.py <slides.json> <input.pptx> <output.pptx>


import argparse
import json
import os
import posixpath
import re
import sys
import zipfile
import xml.etree.ElementTree as ET
from xml.sax.saxutils import quoteattr

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT


# Relationship types whose target part is exclusively "owned" by its referencing part and must
# therefore be deep-copied (rather than shared) whenever the referencing part is duplicated.
# NOTE: deliberately does NOT include the generic RT.SLIDE type - a slide's own .rels can
# contain that when it has a hyperlink jumping to another slide in the deck, and that target
# must stay shared (pointing at the original slide), not be duplicated.
OWNED_RELATIONSHIP_TYPES = frozenset({
    RT.NOTES_SLIDE,
    RT.CHART,
    RT.CHART_COLOR_STYLE,
    RT.CHART_USER_SHAPES,
    RT.OLE_OBJECT,
    RT.PACKAGE,  # embedded chart workbook (.xlsx)
    RT.VML_DRAWING,  # legacy OLE-object icon drawing
    "http://schemas.microsoft.com/office/2011/relationships/chartStyle",
    "http://schemas.microsoft.com/office/2014/relationships/chartEx",  # modern chart types
})

CONTENT_TYPES_NS = "{http://schemas.openxmlformats.org/package/2006/content-types}"
RELATIONSHIPS_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}"

_TRAILING_NUMBER = re.compile(r"^(?P<prefix>.*?)(?P<num>\d+)(?P<ext>\.[^./]+)$")


def build_label_index(prs: Presentation) -> dict[str, str]:
    """
    Scans every slide in a presentation for a shape whose text is wrapped in "[[" and "]]",
    and builds an index mapping each label to the zip-internal part name of its slide.

    Parameters
    ----------
    prs : Presentation
        The python-pptx Presentation to scan

    Returns
    -------
    dict[str, str]
        Mapping of label -> zip-internal slide part name (e.g. "ppt/slides/slide3.xml")
    """
    label_index: dict[str, str] = {}

    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text: str = shape.text
                if text.startswith("[[") and text.endswith("]]"):
                    label = text[2:-2]
                    if label in label_index:
                        print(f"WARNING: label '{label}' found on multiple slides; using the last one",
                              file=sys.stderr)
                    label_index[label] = slide.part.partname.lstrip("/")

    return label_index


def load_labels_from_json(json_path: str) -> list[str]:
    """
    Reads a JSON file (e.g. demo.json) containing a "slides" array and returns the ordered list
    of "layout" values from it - this becomes the list of slide labels to duplicate, in order
    (duplicates allowed - each occurrence produces its own independent copy).

    Parameters
    ----------
    json_path : str
        Path to the JSON file

    Returns
    -------
    list[str]
        Ordered list of "layout" values, one per entry in the "slides" array
    """
    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    if "slides" not in data or not isinstance(data["slides"], list):
        raise ValueError(f"JSON file must contain a 'slides' array: {json_path}")

    labels: list[str] = []
    for i, entry in enumerate(data["slides"]):
        if "layout" not in entry:
            raise ValueError(f"slides[{i}] is missing a 'layout' key in {json_path}")
        labels.append(entry["layout"])
    return labels


def split_numbered_name(partname: str) -> tuple[str, str, str, int]:
    """
    Splits a zip-internal part name into (directory, prefix, extension, trailing number),
    e.g., "ppt/charts/chart3.xml" -> ("ppt/charts/", "chart", ".xml", 3).

    Parameters
    ----------
    partname : str
        Zip-internal part name

    Returns
    -------
    tuple[str, str, str, int]
        (directory including trailing slash, name prefix, extension including leading dot,
        trailing number - 0 if the part name has no trailing digit run)
    """
    directory, basename = posixpath.split(partname)
    if directory:
        directory += "/"
    match = _TRAILING_NUMBER.match(basename)
    if match:
        return directory, match["prefix"], match["ext"], int(match["num"])
    stem, ext = posixpath.splitext(basename)
    return directory, stem, ext, 0


def build_number_counters(namelist: list[str]) -> dict[tuple[str, str, str], int]:
    """
    Seeds a running counter (max number used so far) for every (directory, prefix, extension)
    part-naming pattern found in a zip's namelist, so new parts can be allocated without
    colliding with anything already in the archive.

    Parameters
    ----------
    namelist : list[str]
        Full list of part names in the source .pptx zip archive

    Returns
    -------
    dict[tuple[str, str, str], int]
        Mapping of (directory, prefix, extension) -> the highest number already in use
    """
    counters: dict[tuple[str, str, str], int] = {}
    for name in namelist:
        directory, prefix, ext, num = split_numbered_name(name)
        key = (directory, prefix, ext)
        counters[key] = max(counters.get(key, 0), num)
    return counters


def allocate_new_partname(counters: dict[tuple[str, str, str], int], original_partname: str) -> str:
    """
    Allocates a fresh, collision-free part name for a duplicate of the original_partname, following
    the same directory/prefix/extension naming pattern as the original.

    Parameters
    ----------
    counters : dict[tuple[str, str, str], int]
        Running counters, as built (and mutated in place) by build_number_counters/prior calls
    original_partname : str
        Zip-internal part name being duplicated

    Returns
    -------
    str
        A new, unused zip-internal part name
    """
    directory, prefix, ext, _ = split_numbered_name(original_partname)
    key = (directory, prefix, ext)
    next_num = counters.get(key, 0) + 1
    counters[key] = next_num
    return f"{directory}{prefix}{next_num}{ext}"


def rels_partname_for(partname: str) -> str:
    """Returns the zip-internal path of the .rels file for a given part."""
    directory, name = posixpath.split(partname)
    return posixpath.join(directory, "_rels", name + ".rels")


def resolve_target(part_directory: str, target: str) -> str:
    """Resolves a relationship Target (relative to part_directory) to an absolute zip path."""
    return posixpath.normpath(posixpath.join(part_directory, target))


def relativize(part_directory: str, absolute_target: str) -> str:
    """Expresses absolute_target as a path relative to part_directory."""
    return posixpath.relpath(absolute_target, start=part_directory or ".")


def load_content_types(zip_in: zipfile.ZipFile) -> tuple[dict[str, str], dict[str, str]]:
    """
    Parses [Content_Types].xml (read-only) into its Override and Default entries.

    Parameters
    ----------
    zip_in : zipfile.ZipFile
        The open source .pptx archive

    Returns
    -------
    tuple[dict[str, str], dict[str, str]]
        (overrides: "/ppt/slides/slide1.xml" -> content type,
         defaults: "xml" -> content type)
    """
    root = ET.fromstring(zip_in.read("[Content_Types].xml"))
    overrides: dict[str, str] = {}
    defaults: dict[str, str] = {}
    for child in root:
        if child.tag == f"{CONTENT_TYPES_NS}Override":
            overrides[child.get("PartName")] = child.get("ContentType")
        elif child.tag == f"{CONTENT_TYPES_NS}Default":
            defaults[child.get("Extension")] = child.get("ContentType")
    return overrides, defaults


def content_type_for(overrides: dict[str, str], defaults: dict[str, str],
                      partname: str) -> tuple[str, bool]:
    """
    Looks up the content type for a part, and whether it needs its own explicit Override entry
    (as opposed to being covered by a generic Default-by-extension entry).

    Returns
    -------
    tuple[str, bool]
        (content_type, needs_override)
    """
    key = "/" + partname
    if key in overrides:
        return overrides[key], True
    ext = posixpath.splitext(partname)[1].lstrip(".")
    if ext in defaults:
        return defaults[ext], False
    raise RuntimeError(f"No content type found for part: {partname}")


def duplicate_part(zip_in: zipfile.ZipFile, original_partname: str,
                    content_type_overrides: dict[str, str], content_type_defaults: dict[str, str],
                    counters: dict[tuple[str, str, str], int], staged_parts: dict[str, bytes],
                    new_overrides: list[tuple[str, str]]) -> str:
    """
    Duplicates a part, recursively deep-copying any dependents whose relationship Type is in
    OWNED_RELATIONSHIP_TYPES (e.g. notes, charts). Dependents of any other relationship type
    (images, slide layouts, notes masters, external/same-deck hyperlinks) are left shared - the
    new part's .rels file keeps pointing at the original, unchanged part.

    The part's own XML/binary content is copied byte-for-byte and never rewritten: internal
    relationship-id references (e.g. r:id="rId3") inside that content keep working unmodified
    because the new .rels file keeps the exact same ID values as the original - only Target
    attributes change, and only for owned relationships.

    Parameters
    ----------
    zip_in : zipfile.ZipFile
        The open source .pptx archive
    original_partname : str
        Zip-internal part name to duplicate
    content_type_overrides, content_type_defaults : dict[str, str]
        As returned by load_content_types
    counters : dict[tuple[str, str, str], int]
        Running part-number counters, shared and mutated across the whole run
    staged_parts : dict[str, bytes]
        Accumulator (mutated in place) of new part name -> bytes, shared across the whole run
    new_overrides : list[tuple[str, str]]
        Accumulator (mutated in place) of (new_partname, content_type) pairs that need a new
        Content-Types Override entry

    Returns
    -------
    str
        The new part's zip-internal name
    """
    data = zip_in.read(original_partname)
    new_partname = allocate_new_partname(counters, original_partname)

    content_type, needs_override = content_type_for(content_type_overrides, content_type_defaults,
                                                      original_partname)
    if needs_override:
        new_overrides.append((new_partname, content_type))

    orig_rels_name = rels_partname_for(original_partname)
    if orig_rels_name in zip_in.namelist():
        rels_root = ET.fromstring(zip_in.read(orig_rels_name))
        part_dir = posixpath.dirname(original_partname)
        new_part_dir = posixpath.dirname(new_partname)

        new_rel_lines: list[str] = []
        for rel in rels_root:
            rid = rel.get("Id")
            rtype = rel.get("Type")
            target = rel.get("Target")
            target_mode = rel.get("TargetMode")

            if rtype in OWNED_RELATIONSHIP_TYPES and target_mode != "External":
                dep_abs = resolve_target(part_dir, target)
                new_dep_abs = duplicate_part(zip_in, dep_abs, content_type_overrides,
                                              content_type_defaults, counters, staged_parts,
                                              new_overrides)
                new_target = relativize(new_part_dir, new_dep_abs)
                new_rel_lines.append(
                    f"<Relationship Id={quoteattr(rid)} Type={quoteattr(rtype)} "
                    f"Target={quoteattr(new_target)}/>"
                )
            else:
                mode_attr = f" TargetMode={quoteattr(target_mode)}" if target_mode else ""
                new_rel_lines.append(
                    f"<Relationship Id={quoteattr(rid)} Type={quoteattr(rtype)} "
                    f"Target={quoteattr(target)}{mode_attr}/>"
                )

        new_rels_xml = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            + "".join(new_rel_lines) + "</Relationships>"
        ).encode("utf-8")
        staged_parts[rels_partname_for(new_partname)] = new_rels_xml

    staged_parts[new_partname] = data
    return new_partname


def duplicate_labeled_slides(zip_in: zipfile.ZipFile, labels_to_duplicate: list[str],
                              label_index: dict[str, str], content_type_overrides: dict[str, str],
                              content_type_defaults: dict[str, str],
                              counters: dict[tuple[str, str, str], int]
                              ) -> tuple[list[str], dict[str, bytes], list[tuple[str, str]]]:
    """
    Duplicates every slide named in labels_to_duplicate, in order. If the same label appears
    more than once, each occurrence produces its own fully independent set of new parts (there
    is no memoization of duplicate_part across these top-level calls).

    Returns
    -------
    tuple[list[str], dict[str, bytes], list[tuple[str, str]]]
        (new slide part names in duplication order, staged new parts, new Content-Types overrides)
    """
    staged_parts: dict[str, bytes] = {}
    new_overrides: list[tuple[str, str]] = []
    new_slide_partnames: list[str] = []

    for label in labels_to_duplicate:
        original_slide_partname = label_index[label]
        new_partname = duplicate_part(zip_in, original_slide_partname, content_type_overrides,
                                       content_type_defaults, counters, staged_parts, new_overrides)
        new_slide_partnames.append(new_partname)

    return new_slide_partnames, staged_parts, new_overrides


def patch_presentation_rels(rels_bytes: bytes, new_slide_partnames: list[str]) -> tuple[bytes, list[str]]:
    """
    Inserts one new slide Relationship per duplicated slide into ppt/_rels/presentation.xml.rels.

    Returns
    -------
    tuple[bytes, list[str]]
        (patched file bytes, the rId assigned to each new slide, in the same order as
        new_slide_partnames)
    """
    text = rels_bytes.decode("utf-8")
    existing_ids = [int(n) for n in re.findall(r'Id="rId(\d+)"', text)]
    next_id = max(existing_ids, default=0) + 1

    assigned_rids: list[str] = []
    new_lines: list[str] = []
    for new_partname in new_slide_partnames:
        rid = f"rId{next_id}"
        next_id += 1
        assigned_rids.append(rid)
        target = posixpath.relpath(new_partname, start="ppt")
        new_lines.append(
            f'<Relationship Id={quoteattr(rid)} Type={quoteattr(RT.SLIDE)} '
            f'Target={quoteattr(target)}/>'
        )

    if text.count("</Relationships>") != 1:
        raise RuntimeError("Could not find a single </Relationships> closing tag in "
                            "ppt/_rels/presentation.xml.rels")
    patched = text.replace("</Relationships>", "".join(new_lines) + "</Relationships>", 1)
    return patched.encode("utf-8"), assigned_rids


def patch_presentation_xml(pres_bytes: bytes, assigned_rids: list[str]) -> bytes:
    """
    Inserts a new <p:sldId> element for each newly duplicated slide at the very front of
    <p:sldIdLst> in ppt/presentation.xml, in the same order as assigned_rids, so the duplicates
    appear before all original slides once the deck is opened.

    Returns
    -------
    bytes
        Patched file bytes
    """
    text = pres_bytes.decode("utf-8")

    prefix_match = re.search(r"<(\w+):sldIdLst>", text)
    if not prefix_match:
        raise RuntimeError("Could not find <p:sldIdLst> opening tag in ppt/presentation.xml")
    prefix = prefix_match.group(1)

    existing_ids = [int(n) for n in re.findall(rf'<{prefix}:sldId\s+id="(\d+)"', text)]
    next_id = max(max(existing_ids, default=255), 255) + 1

    new_elems: list[str] = []
    for rid in assigned_rids:
        new_elems.append(f'<{prefix}:sldId id="{next_id}" r:id={quoteattr(rid)}/>')
        next_id += 1

    open_tag = f"<{prefix}:sldIdLst>"
    if text.count(open_tag) != 1:
        raise RuntimeError(f"Could not find a single {open_tag} opening tag in ppt/presentation.xml")
    patched = text.replace(open_tag, open_tag + "".join(new_elems), 1)
    return patched.encode("utf-8")


def patch_content_types(ct_bytes: bytes, new_overrides: list[tuple[str, str]]) -> bytes:
    """
    Inserts a new <Override> entry for each newly duplicated part that needs one into
    [Content_Types].xml.

    Returns
    -------
    bytes
        Patched file bytes
    """
    text = ct_bytes.decode("utf-8")
    new_lines = [
        f"<Override PartName={quoteattr('/' + partname)} ContentType={quoteattr(content_type)}/>"
        for partname, content_type in new_overrides
    ]
    if text.count("</Types>") != 1:
        raise RuntimeError("Could not find a single </Types> closing tag in [Content_Types].xml")
    patched = text.replace("</Types>", "".join(new_lines) + "</Types>", 1)
    return patched.encode("utf-8")


def write_output_pptx(zip_in: zipfile.ZipFile, output_path: str, patched: dict[str, bytes],
                       staged_parts: dict[str, bytes]) -> None:
    """
    Writes the new .pptx: every original entry is copied through unchanged except the patched
    bookkeeping files, plus all newly staged (duplicated) parts.

    Parameters
    ----------
    zip_in : zipfile.ZipFile
        The open source .pptx archive
    output_path : str
        Path to write the new .pptx to
    patched : dict[str, bytes]
        Bookkeeping files that must be replaced (by zip-internal name)
    staged_parts : dict[str, bytes]
        Newly duplicated parts to add (by zip-internal name)
    """
    with zipfile.ZipFile(output_path, "w", zipfile.ZIP_DEFLATED) as zip_out:
        zip_out.writestr("[Content_Types].xml", patched["[Content_Types].xml"])
        for info in zip_in.infolist():
            if info.filename == "[Content_Types].xml":
                continue
            if info.filename in patched:
                zip_out.writestr(info, patched[info.filename])
            else:
                zip_out.writestr(info, zip_in.read(info.filename))
        for new_partname, part_data in staged_parts.items():
            zip_out.writestr(new_partname, part_data)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Duplicate labeled slides in a .pptx and prepend them to the deck."
    )
    parser.add_argument("json_file", help="Path to a JSON file listing slides by layout label (e.g. demo.json)")
    parser.add_argument("input_pptx", help="Path to the source .pptx file")
    parser.add_argument("output_pptx", help="Path to save the new .pptx file to")
    args = parser.parse_args()

    if not os.path.isfile(args.json_file):
        sys.exit(f"ERROR: JSON file not found: {args.json_file}")
    if not os.path.isfile(args.input_pptx):
        sys.exit(f"ERROR: Input .pptx file not found: {args.input_pptx}")
    if os.path.isfile(args.output_pptx):
        sys.exit("ERROR Output .pptx file already exists - please delete it and try again")

    try:
        labels_to_duplicate = load_labels_from_json(args.json_file)
    except (json.JSONDecodeError, ValueError, OSError) as e:
        sys.exit(f"ERROR: could not read JSON file {args.json_file}: {e}")

    prs = Presentation(args.input_pptx)
    label_index = build_label_index(prs)

    missing_labels: list = [label for label in labels_to_duplicate if label not in label_index]
    if missing_labels:
        sys.exit(f"ERROR: label(s) not found in presentation: {', '.join(missing_labels)}")

    with zipfile.ZipFile(args.input_pptx, "r") as zip_in:
        content_type_overrides, content_type_defaults = load_content_types(zip_in)
        counters = build_number_counters(zip_in.namelist())

        new_slide_partnames, staged_parts, new_overrides = duplicate_labeled_slides(
            zip_in, labels_to_duplicate, label_index, content_type_overrides,
            content_type_defaults, counters
        )

        patched_rels, assigned_rids = patch_presentation_rels(
            zip_in.read("ppt/_rels/presentation.xml.rels"), new_slide_partnames
        )
        patched_pres = patch_presentation_xml(zip_in.read("ppt/presentation.xml"), assigned_rids)
        patched_ct = patch_content_types(zip_in.read("[Content_Types].xml"), new_overrides)

        patched = {
            "ppt/_rels/presentation.xml.rels": patched_rels,
            "ppt/presentation.xml": patched_pres,
            "[Content_Types].xml": patched_ct,
        }

        all_new_xml = dict(patched)
        for name, blob in staged_parts.items():
            if name.endswith(".xml") or name.endswith(".rels"):
                all_new_xml[name] = blob
        for name, blob in all_new_xml.items():
            try:
                ET.fromstring(blob)
            except ET.ParseError as e:
                sys.exit(f"ERROR: generated invalid XML for {name}: {e}")

        write_output_pptx(zip_in, args.output_pptx, patched, staged_parts)

    bad_entry = zipfile.ZipFile(args.output_pptx).testzip()
    if bad_entry is not None:
        sys.exit(f"ERROR: output .pptx failed integrity check at: {bad_entry}")

    verify_prs = Presentation(args.output_pptx)
    expected_count = len(prs.slides) + len(labels_to_duplicate)
    actual_count = len(verify_prs.slides)
    if actual_count != expected_count:
        sys.exit(f"ERROR: output .pptx has {actual_count} slide(s), expected {expected_count}")

    print(f"Duplicated {len(labels_to_duplicate)} slide(s) and saved to: {args.output_pptx}")
    print(f"Output deck now has {actual_count} slide(s) (was {len(prs.slides)})")


if __name__ == "__main__":
    main()
