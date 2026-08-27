from pptx import Presentation


def get_slide_layouts(prs):

    found_layouts: dict = dict()

    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text: str = shape.text
                if text.startswith("[[") and text.endswith("]]"):
                    key = text[2:-2]
                    found_layouts[key] = slide

    return found_layouts


# def duplicate_slide_and_place_at_start(prs, layouts: dict, key: str):
#     if key not in layouts:
#         print(f"Slide layout '{key}' not available.")
#         return
#
#     # Duplicate the slide
#     new_slide = layouts[key]._slide.duplicate()
#     prs.slides._spTree.insert(0, new_slide._spTree)
#
#     # Return the updated prs
#     return prs


new_prs = Presentation("demo.pptx")


